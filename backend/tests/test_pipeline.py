import io
import json
import zipfile

import pytest

from app.core import store as s
from app.ingestion import indexing as ix
from app.ingestion.extraction import Extractor
from app.integrations import providers as p
from app.main import register

def test_source_inventory_and_provenance(indexed):
    doc, index = indexed
    assert doc["active_index"] == index["id"]
    assert {e["kind"] for e in s.get("extractions", index["extraction_id"])["elements"]} >= {
        "section_header",
        "text",
    }
    assert all(c["element_ids"] and c["offsets"] for c in index["chunks"])

def test_search_and_index_isolation(indexed, fake_vectors):
    doc, index = indexed
    hits = ix.search(index, "solar", p.Budget("search"), 2)
    assert "Solar" in hits[0]["text"]
    other = register("private.txt", b"Unrelated private material")
    ix.extract_document(other["id"], index["config"], "e2")
    ix.build_index(other["id"], index["config"], "i2")
    assert all(
        "private" not in h["text"].lower()
        for h in ix.search(index, "solar", p.Budget("search"), 20)
    )

def test_invalid_overlap():
    with pytest.raises(ValueError, match="Overlap"):
        ix.split_text("abcdefgh", 4, 4)

def test_all_text_covered():
    text = "αβγ\n" + ("solar batteries. " * 500)
    pieces = ix.split_text(text, 200, 20, True)
    covered = set()
    for part, start, end in pieces:
        assert text[start:end] == part
        covered.update(range(start, end))
    assert covered == set(range(len(text)))

def test_invalid_agentic_plan_rejected(monkeypatch, cfg):
    cfg.strategy = "agentic"
    elems = [{"id": "e1", "kind": "text", "section": "", "text": "real evidence"}]
    monkeypatch.setattr(
        p, "generate", lambda *a, **k: {"groups": [{"unit_ids": ["invented"], "title": "bad"}]}
    )
    with pytest.raises(ValueError, match="omitted"):
        ix.make_chunks(elems, cfg, {"id": "dummy"}, p.Budget("plan"))

def test_agentic_chunks_preserve_verbatim(monkeypatch, cfg):
    cfg.strategy = "agentic"
    elems = [
        {"id": "e1", "kind": "text", "section": "", "text": "First verified fact."},
        {"id": "e2", "kind": "text", "section": "", "text": "Second fact."},
    ]
    monkeypatch.setattr(
        p, "generate", lambda *a, **k: {"groups": [{"unit_ids": ["u1", "u2"], "title": "Topic"}]}
    )
    out = ix.make_chunks(elems, cfg, {"id": "dummy"}, p.Budget("plan"))
    assert out[0]["text"] == "First verified fact.\n\nSecond fact."
    assert out[0]["element_ids"] == ["e1", "e2"]

def test_failed_index_keeps_previous(indexed, monkeypatch):
    doc, index = indexed
    monkeypatch.setattr(
        ix, "vectors", lambda *args: (_ for _ in ()).throw(ValueError("embedding unavailable"))
    )
    with pytest.raises(ValueError, match="unavailable"):
        ix.build_index(doc["id"], index["config"], "fail")
    assert s.get("documents", doc["id"])["active_index"] == index["id"]

def test_multimodal_requires_image_encoder(indexed):
    doc, index = indexed
    cfg = {**index["config"], "mode": "multimodal", "embedding": "provider-text"}
    with pytest.raises(ValueError, match="True multimodal"):
        ix.build_index(doc["id"], cfg, "mode")

def test_preprocessing_preserves_inspector(cfg):
    cfg.mask_pii = True
    cfg.strip_headers = True
    elems = [
        {"id": "e1", "kind": "text", "text": "Contact person@example.org", "section": ""},
        {"id": "e2", "kind": "header", "text": "Repeated", "section": ""},
    ]
    out = ix.preprocess(elems, cfg)
    assert len(out) == 1 and "[EMAIL]" in out[0]["prepared_text"]
    assert elems[0]["text"] == "Contact person@example.org"

def test_excel_formulas_images_and_rows(cfg):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Revenue"])
    ws.append(["Jan", 200])
    ws.append(["Feb", 300])
    ws["B4"] = "=SUM(B2:B3)"
    b = io.BytesIO()
    wb.save(b)
    d = register("sales.xlsx", b.getvalue())
    elements, warnings = Extractor(d, cfg, None, p.Budget("xlsx")).run()
    assert any(e["kind"] == "table" and "Revenue" in e["text"] for e in elements)
    assert any(e["kind"] == "formula" and "=SUM(B2:B3)" in e["text"] for e in elements)
    assert any("never executed" in w for w in warnings)

def test_docx_headers_tables_and_multiple_parts(cfg):
    from docx import Document

    d = Document()
    d.add_heading("Report", 0)
    d.add_paragraph("Evidence")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Name"
    t.cell(1, 0).text = "Solar"
    d.sections[0].header.paragraphs[0].text = "Company header"
    b = io.BytesIO()
    d.save(b)
    dest = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(b.getvalue())) as source, zipfile.ZipFile(dest, "w") as target:
        for n in source.namelist():
            target.writestr(n, source.read(n))
        target.writestr(
            "word/document2.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Secondary document content</w:t></w:r></w:p></w:body></w:document>',
        )
    source = register("report.docx", dest.getvalue())
    elements, _ = Extractor(source, cfg, None, p.Budget("word")).run()
    assert any(e["kind"] == "table" and "Solar" in e["text"] for e in elements)
    assert any(e["kind"] == "header" and "Company" in e["text"] for e in elements)
    assert any("Secondary document" in e["text"] for e in elements)

def test_html_never_runs_scripts(cfg):
    d = register(
        "page.html",
        b'<h1>Solar</h1><p>Real source</p><script>secret()</script><img src="http://127.0.0.1/secret" alt="Solar array"><table><tr><td>12</td></tr></table>',
    )
    elements, _ = Extractor(d, cfg, None, p.Budget("html")).run()
    assert not any("secret()" in e["text"] for e in elements)
    assert any(e["kind"] == "image" and e["text"] == "Solar array" for e in elements)

def test_notebook_is_data(cfg):
    notebook = {
        "cells": [
            {
                "cell_type": "code",
                "source": ['raise RuntimeError("never run")'],
                "outputs": [{"data": {"text/plain": ["Saved result"]}}],
            }
        ]
    }
    d = register("code.ipynb", json.dumps(notebook).encode())
    elements, _ = Extractor(d, cfg, None, p.Budget("notebook")).run()
    assert any(e["kind"] == "code" for e in elements)
    assert any(e["kind"] == "output" and "Saved result" in e["text"] for e in elements)

def test_pdf_native_assets(cfg):
    import pymupdf as fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((40, 60), "Solar report", fontsize=22)
    page.insert_text((40, 100), "Panels convert sunlight to electricity.")
    d = register("source.pdf", pdf.tobytes())
    elements, _ = Extractor(d, cfg, None, p.Budget("pdf")).run()
    assert any(e["kind"] == "page_preview" and s.asset_path(e["asset"]).exists() for e in elements)
    assert any("Panels convert" in e["text"] and e["page"] == 1 for e in elements)

def test_visual_units_never_hide_second_asset(monkeypatch, cfg):
    cfg.strategy = "agentic"
    es = [
        {"id": "e1", "kind": "image", "text": "red square", "asset": "assets/a.png", "section": ""},
        {
            "id": "e2",
            "kind": "image",
            "text": "blue circle",
            "asset": "assets/b.png",
            "section": "",
        },
    ]
    monkeypatch.setattr(
        p, "generate", lambda *a, **k: {"groups": [{"unit_ids": ["u1", "u2"], "title": "shapes"}]}
    )
    chunks = ix.make_chunks(es, cfg, {"id": "dummy"}, p.Budget("visual-plan"))
    assert len(chunks) == 2 and chunks[0]["assets"] != chunks[1]["assets"]

def test_excel_chart_has_rendered_data_asset(cfg):
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Revenue"])
    ws.append(["Jan", 20])
    ws.append(["Feb", 30])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=3))
    ws.add_chart(chart, "D2")
    b = io.BytesIO()
    wb.save(b)
    doc = register("chart.xlsx", b.getvalue())
    elements, _ = Extractor(doc, cfg, None, p.Budget("chart")).run()
    chart = next(e for e in elements if e["kind"] == "chart")
    assert chart["asset"] and chart["data_asset"] and chart["preview_rendered"]
    assert "20.0" in chart["text"] and "30.0" in chart["text"]

def test_cancelled_index_never_activates(indexed, monkeypatch):
    doc, index = indexed
    original = ix.vectors
    s.put("jobs", "cancel-build", {"kind": "indexing", "status": "running", "cancel": False})

    def cancel_after_vectors(*args, **kwargs):
        result = original(*args, **kwargs)
        s.patch("jobs", "cancel-build", cancel=True)
        return result

    monkeypatch.setattr(ix, "vectors", cancel_after_vectors)
    with pytest.raises(ValueError, match="Cancelled"):
        ix.build_index(doc["id"], index["config"], "cancel-build")
    assert s.get("documents", doc["id"])["active_index"] == index["id"]
    s.delete("jobs", "cancel-build")

def test_html_table_has_rows_and_visual_preview(cfg):
    d = register(
        "table.html",
        b"<table><tr><th>Name</th><th>Value</th></tr><tr><td>Solar</td><td>12</td></tr></table>",
    )
    elements, _ = Extractor(d, cfg, None, p.Budget("html-table")).run()
    table = next(e for e in elements if e["kind"] == "table")
    assert table["rows"] == [["Name", "Value"], ["Solar", "12"]]
    assert table["asset"] and "Solar | 12" in table["text"]

def test_legacy_word_converts_locally(cfg, tmp_path):
    import shutil
    import subprocess

    if not shutil.which("textutil"):
        pytest.skip("macOS textutil not installed")
    text = tmp_path / "legacy.txt"
    text.write_text("Legacy Word evidence about solar power.")
    output = tmp_path / "legacy.doc"
    subprocess.run(["textutil", "-convert", "doc", "-output", str(output), str(text)], check=True)
    doc = register("legacy.doc", output.read_bytes())
    elements, warnings = Extractor(doc, cfg, None, p.Budget("legacy-doc")).run()
    assert any("Legacy Word evidence" in e["text"] for e in elements)
    assert any("converted locally" in w for w in warnings)
